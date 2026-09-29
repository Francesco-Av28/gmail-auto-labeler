# -*- coding: utf-8 -*-
"""
RIORDINO EMAIL - riordina un archivio Gmail con etichette, tramite Gmail API.

Fasi (ognuna e' un comando separato; solo --esegui modifica la casella):
  python riordino_email.py --login    autorizza l'accesso (apre il browser, una volta sola)
  python riordino_email.py --estrai   scarica i METADATI (mittente, destinatari, oggetto, anteprima, data,
                                      etichette) dei messaggi in messaggi.csv. Nessuna modifica.
  python riordino_email.py --piano    PROVA A SECCO: calcola la destinazione di ogni conversazione ->
                                      piano.csv + piano_riordino.xlsx. Lavora offline, nessuna modifica.
  python riordino_email.py --esegui   applica il piano (crea le etichette, etichetta, toglie dalla Posta
                                      in arrivo se previsto). Chiede conferma scrivendo SI.

Regole (tutte in configurazione.json):
  - l'anno di una conversazione e' quello del suo ULTIMO messaggio;
  - "fasce_anni" decide per ogni periodo: archivio per anno, classificazione, oppure nessuna azione;
  - classificazione: 1) rubrica per indirizzo email, 2) nome del contatto in oggetto/anteprima/nomi,
    3) regole regex su mittente/oggetto, 4) autoinviate / caselle personali / altro.
Nulla viene mai cestinato o eliminato: si aggiungono etichette e, se previsto, si toglie INBOX.
"""
import argparse
import csv
import json
import os
import re
import sys
import time
import unicodedata
from collections import defaultdict, Counter
from datetime import datetime, timezone
from email.utils import getaddresses

BASE = os.path.dirname(os.path.abspath(__file__))
CONFIG = os.path.join(BASE, "configurazione.json")
CREDENZIALI = os.path.join(BASE, "credentials.json")
TOKEN = os.path.join(BASE, "token.json")
STATO = os.path.join(BASE, "stato.json")              # indirizzo della casella + etichette esistenti
MESSAGGI_CSV = os.path.join(BASE, "messaggi.csv")
PIANO_CSV = os.path.join(BASE, "piano.csv")
PIANO_XLSX = os.path.join(BASE, "piano_riordino.xlsx")
LOG = os.path.join(BASE, "esecuzione.log")

SCOPES = ["https://www.googleapis.com/auth/gmail.modify", "https://www.googleapis.com/auth/gmail.labels"]
FORME_GIURIDICHE = ["SOC COOP", "SOCIETA COOPERATIVA", "SRLS", "SRL", "SAS", "SNC", "SPA", "STP", "ASD"]
PAROLE_VUOTE = {"DI", "DEI", "DEL", "DELLA", "DELLE", "LA", "IL", "LE", "E", "AND", "SOC"}


# ----------------------------------------------------------------------------- utilita'
def log(msg):
    riga = f"{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}"
    print(riga)
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(riga + "\n")


def carica_config(percorso=CONFIG):
    with open(percorso, encoding="utf-8") as f:
        cfg = {k: v for k, v in json.load(f).items() if not k.startswith("_")}
    cfg["_regole"] = [(r["etichetta"],
                       re.compile(r["mittente"], re.I) if r.get("mittente") else None,
                       re.compile(r["oggetto"], re.I) if r.get("oggetto") else None)
                      for r in cfg.get("regole", [])]
    cfg["_personali"] = re.compile(cfg.get("regex_caselle_personali") or r"(?!)", re.I)
    return cfg


def carica_stato():
    if os.path.exists(STATO):
        with open(STATO, encoding="utf-8") as f:
            return json.load(f)
    return {"indirizzo": "", "etichette": {}}


def salva_stato(stato):
    with open(STATO, "w", encoding="utf-8") as f:
        json.dump(stato, f, ensure_ascii=False, indent=1)


def normalizza(s):
    """Maiuscolo, senza accenti e punteggiatura; sigle puntate compattate ('S.r.l.' -> 'SRL')."""
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.upper().replace("'", "").replace("’", "")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    s = " " + re.sub(r"\s+", " ", s).strip() + " "
    s = re.sub(r"(?<= )([A-Z](?: [A-Z])+)(?= )", lambda m: m.group(1).replace(" ", ""), s)
    return s.strip()


def nucleo_nome(nome):
    """'Rossi Mario S.r.l.' -> 'ROSSI MARIO' (senza forma giuridica)."""
    n = normalizza(nome)
    for forma in FORME_GIURIDICHE:
        n = re.sub(rf"\b{forma}\b", " ", n)
    return re.sub(r"\s+", " ", n).strip()


def indirizzi(campo):
    return [(n.strip(), a.strip().lower()) for n, a in getaddresses([campo or ""]) if a]


def carica_rubrica(cfg):
    """{indirizzo email: nome contatto} dal CSV CONTATTO;EMAIL (facoltativo)."""
    rub = {}
    nome_file = cfg.get("rubrica_csv")
    if not nome_file:
        return rub
    percorso = nome_file if os.path.isabs(nome_file) else os.path.join(BASE, nome_file)
    if not os.path.exists(percorso):
        log(f"Rubrica non trovata ({percorso}): si prosegue senza")
        return rub
    with open(percorso, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            nome = (r.get("CONTATTO") or "").strip()
            for em in re.split(r"[;, ]+", r.get("EMAIL") or ""):
                em = em.strip().lower()
                if nome and "@" in em:
                    rub[em] = nome
    return rub


# ----------------------------------------------------------------------------- classificazione (offline)
def fascia_di(anno, cfg):
    for f in cfg["fasce_anni"]:
        if f["da"] <= anno <= f["a"]:
            return f
    return None


def prepara_nomi(rubrica, cfg, etichette_esistenti=()):
    """Nomi da cercare nei testi: contatti della rubrica + etichette gia' presenti sotto la madre rubrica.
    Ritorna [(etichetta, nucleo, rischioso)]; i nomi corti richiedono la forma giuridica accanto."""
    madre = cfg["etichetta_madre_rubrica"]
    nomi = {f"{madre}/{n}" for n in rubrica.values()}
    nomi |= {e for e in etichette_esistenti if e.startswith(madre + "/")}
    elenco = []
    for etichetta in sorted(nomi):
        nuc = nucleo_nome(etichetta.split("/", 1)[1])
        if not nuc:
            continue
        rischioso = len(nuc.replace(" ", "")) <= 4 or (len(nuc.split()) == 1 and len(nuc) <= 6)
        elenco.append((etichetta, nuc, rischioso))
    return elenco


def classifica_conversazione(msgs, cfg, rubrica, nomi, mio_indirizzo=""):
    """Destinazione di una conversazione (lista di messaggi-dict ordinati per data).
    Ritorna (destinazione, metodo, dettaglio); destinazione '' = non toccare."""
    ultimo = msgs[-1]
    anno = int(ultimo["data"][:4])
    fascia = fascia_di(anno, cfg)
    if fascia is None or fascia["azione"] == "nessuna":
        return "", "nessuna azione", f"anno {anno}"
    if fascia["azione"] == "archivio_per_anno":
        return f"{fascia['prefisso']}/{anno}", "anno", ""

    madre = cfg["etichetta_madre_rubrica"]
    # 1) rubrica per indirizzo
    trovati = Counter()
    for m in msgs:
        for campo in ("from", "to", "cc"):
            for _, a in indirizzi(m.get(campo)):
                if a in rubrica:
                    trovati[f"{madre}/{rubrica[a]}"] += 1
    if trovati:
        return trovati.most_common(1)[0][0], "rubrica", "indirizzo in rubrica"

    # 2) nome del contatto nei testi
    if cfg.get("cerca_nomi_rubrica_nei_testi", True) and nomi:
        punteggi = Counter()
        for m in msgs:
            testi = [normalizza(m.get("subject")), normalizza(m.get("snippet"))]
            testi += [normalizza(n) for c in ("from", "to", "cc") for n, a in indirizzi(m.get(c))
                      if a != mio_indirizzo]
            for etichetta, nuc, rischioso in nomi:
                for t in testi:
                    if not t:
                        continue
                    if rischioso:
                        if re.search(rf"\b{re.escape(nuc)}\b\s*(SRLS|SRL|SAS|SNC|SPA|ASD|STP)\b", t):
                            punteggi[etichetta] += 3
                    elif re.search(rf"\b{re.escape(nuc)}\b", t):
                        punteggi[etichetta] += 3
        if punteggi:
            top = punteggi.most_common()
            if not (len(top) > 1 and top[0][1] == top[1][1]):
                return top[0][0], "nome", top[0][0]

    # 3) regole su mittente / oggetto (dell'ultimo messaggio)
    madre_regole = cfg["etichetta_madre_regole"]
    mittente, oggetto = ultimo.get("from") or "", ultimo.get("subject") or ""
    for nome, rm, ro in cfg["_regole"]:
        if (rm and rm.search(mittente)) or (ro and ro.search(oggetto)):
            return f"{madre_regole}/{nome}", "regola", nome

    # 4) finali
    indirizzi_mitt = [a for _, a in indirizzi(mittente)]
    if mio_indirizzo and mio_indirizzo in indirizzi_mitt:
        return f"{madre_regole}/{cfg['etichetta_autoinviate']}", "autoinviata", ""
    if cfg["_personali"].search(mittente):
        return f"{madre_regole}/{cfg['etichetta_contatti_personali']}", "casella personale", ""
    return f"{madre_regole}/{cfg['etichetta_altro']}", "altro", ""


def calcola_piano(messaggi, cfg, rubrica, etichette_esistenti=None, mio_indirizzo=""):
    """messaggi: righe di messaggi.csv. Ritorna la lista del piano, una riga per conversazione."""
    etichette_esistenti = etichette_esistenti or {}
    id2nome = {v: k for k, v in etichette_esistenti.items()}
    nomi = prepara_nomi(rubrica, cfg, etichette_esistenti)
    thread = defaultdict(list)
    for r in messaggi:
        thread[r["threadId"]].append(r)
    piano = []
    for tid, msgs in thread.items():
        msgs.sort(key=lambda m: int(m["internalDate"]))
        ultimo = msgs[-1]
        dest, metodo, dett = classifica_conversazione(msgs, cfg, rubrica, nomi, mio_indirizzo)
        etichette_attuali = sorted({id2nome.get(l, l) for m in msgs for l in (m.get("labelIds") or "").split("|")
                                    if l.startswith("Label_")})
        piano.append({
            "threadId": tid, "anno": int(ultimo["data"][:4]), "n_messaggi": len(msgs),
            "in_inbox": "SI" if any("INBOX" in (m.get("labelIds") or "").split("|") for m in msgs) else "no",
            "data_ultimo": ultimo["data"], "da": (ultimo.get("from") or "")[:80], "a": (ultimo.get("to") or "")[:80],
            "oggetto": (ultimo.get("subject") or "")[:120], "etichette_attuali": " | ".join(etichette_attuali),
            "destinazione": dest, "metodo": metodo, "dettaglio": dett[:150],
            "messageIds": "|".join(m["id"] for m in msgs),
        })
    piano.sort(key=lambda p: (p["anno"], p["destinazione"], p["data_ultimo"]))
    return piano


# ----------------------------------------------------------------------------- Gmail API
def servizio():
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from google.auth.transport.requests import Request
    from googleapiclient.discovery import build
    creds = None
    if os.path.exists(TOKEN):
        creds = Credentials.from_authorized_user_file(TOKEN, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists(CREDENZIALI):
                sys.exit(f"Manca il file {CREDENZIALI}: crearlo sulla console Google Cloud (vedi LEGGIMI.txt).")
            flow = InstalledAppFlow.from_client_secrets_file(CREDENZIALI, SCOPES)
            creds = flow.run_local_server(port=0, prompt="consent")
        with open(TOKEN, "w") as f:
            f.write(creds.to_json())
    return build("gmail", "v1", credentials=creds, cache_discovery=False)


def con_ritentativi(fn, tentativi=6):
    for i in range(tentativi):
        try:
            return fn()
        except Exception as e:  # noqa
            attesa = 2 ** i
            log(f"  errore API ({e.__class__.__name__}: {str(e)[:120]}), riprovo tra {attesa}s")
            time.sleep(attesa)
    raise RuntimeError("troppi errori consecutivi")


def aggiorna_stato_da_server(svc):
    stato = carica_stato()
    stato["indirizzo"] = svc.users().getProfile(userId="me").execute()["emailAddress"].lower()
    etichette = con_ritentativi(lambda: svc.users().labels().list(userId="me").execute())["labels"]
    stato["etichette"] = {l["name"]: l["id"] for l in etichette}
    salva_stato(stato)
    return stato


def cmd_login():
    svc = servizio()
    prof = svc.users().getProfile(userId="me").execute()
    aggiorna_stato_da_server(svc)
    log(f"Autorizzato - messaggi totali {prof['messagesTotal']}, conversazioni {prof['threadsTotal']}")


CAMPI = ["id", "threadId", "internalDate", "data", "labelIds", "from", "to", "cc", "subject", "snippet"]


def cmd_estrai(query):
    """Scarica i metadati dei messaggi in messaggi.csv (sola lettura). Riprende se interrotto."""
    svc = servizio()
    aggiorna_stato_da_server(svc)
    gia = set()
    if os.path.exists(MESSAGGI_CSV):
        with open(MESSAGGI_CSV, encoding="utf-8", newline="") as f:
            gia = {r["id"] for r in csv.DictReader(f)}
        log(f"Ripresa: {len(gia)} messaggi gia' scaricati")
    ids, token = [], None
    while True:
        r = con_ritentativi(lambda: svc.users().messages().list(
            userId="me", q=query, maxResults=500, pageToken=token, includeSpamTrash=False).execute())
        ids += [m["id"] for m in r.get("messages", [])]
        token = r.get("nextPageToken")
        print(f"\r  elencati {len(ids)} messaggi", end="")
        if not token:
            break
    print()
    da_fare = [i for i in ids if i not in gia]
    log(f"Query [{query}]: messaggi {len(ids)}, da scaricare {len(da_fare)}")

    nuovo = not os.path.exists(MESSAGGI_CSV)
    fout = open(MESSAGGI_CSV, "a", encoding="utf-8", newline="")
    w = csv.DictWriter(fout, fieldnames=CAMPI)
    if nuovo:
        w.writeheader()
    righe, falliti = [], []

    def cb(rid, resp, exc):
        if exc is not None:
            falliti.append(rid)
            return
        h = {x["name"].lower(): x["value"] for x in resp.get("payload", {}).get("headers", [])}
        ts = int(resp["internalDate"]) // 1000
        righe.append({"id": resp["id"], "threadId": resp["threadId"], "internalDate": resp["internalDate"],
                      "data": datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M"),
                      "labelIds": "|".join(resp.get("labelIds", [])),
                      "from": h.get("from", ""), "to": h.get("to", ""), "cc": h.get("cc", ""),
                      "subject": h.get("subject", ""), "snippet": resp.get("snippet", "")})

    # Quota Gmail: ~15.000 unita'/minuto per utente, messages.get = 5 unita'. Blocchi da 25 ogni 2 s
    # (~750/min) e, se Google rifiuta per quota, pausa crescente e nuovo tentativo.
    fatti, coda, giro_errori = 0, list(da_fare), 0
    while coda:
        blocco, coda = coda[:25], coda[25:]
        falliti.clear()
        batch = svc.new_batch_http_request(callback=cb)
        for mid in blocco:
            batch.add(svc.users().messages().get(userId="me", id=mid, format="metadata",
                                                 metadataHeaders=["From", "To", "Cc", "Subject"]), request_id=mid)
        con_ritentativi(batch.execute)
        w.writerows(righe)
        fout.flush()
        fatti += len(righe)
        righe.clear()
        if falliti:
            giro_errori += 1
            coda = list(falliti) + coda
            attesa = min(60, 10 * giro_errori)
            log(f"  quota superata su {len(falliti)} messaggi: pausa {attesa}s")
            time.sleep(attesa)
        else:
            giro_errori = 0
            time.sleep(2.0)
        print(f"\r  scaricati {fatti}/{len(da_fare)}", end="")
    print()
    fout.close()
    log(f"Estrazione completata: {fatti} nuovi messaggi in messaggi.csv")


# ----------------------------------------------------------------------------- --piano (prova a secco)
def cmd_piano():
    cfg = carica_config()
    if not os.path.exists(MESSAGGI_CSV):
        sys.exit("Manca messaggi.csv: lanciare prima --estrai.")
    stato = carica_stato()
    rubrica = carica_rubrica(cfg)
    with open(MESSAGGI_CSV, encoding="utf-8", newline="") as f:
        messaggi = list(csv.DictReader(f))
    piano = calcola_piano(messaggi, cfg, rubrica, stato.get("etichette"), stato.get("indirizzo", ""))
    if not piano:
        sys.exit("Nessun messaggio in messaggi.csv.")
    with open(PIANO_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(piano[0].keys()))
        w.writeheader()
        w.writerows(piano)
    log(f"PROVA A SECCO - nessuna modifica alla casella. Conversazioni: {len(piano)}")
    for d, n in sorted(Counter(p["destinazione"] or "(non toccate)" for p in piano).items()):
        log(f"  {n:>7}  {d}")
    scrivi_excel(piano)
    log(f"Piano salvato: {PIANO_CSV} e {PIANO_XLSX}. Controllarlo, poi eventualmente --esegui.")


def scrivi_excel(piano):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Riepilogo"
    ws.append(["Destinazione", "Conversazioni"])
    for d, n in sorted(Counter(p["destinazione"] or "(non toccate)" for p in piano).items()):
        ws.append([d, n])
    ws.column_dimensions["A"].width = 50
    for c in ws[1]:
        c.font = Font(bold=True)
    colonne = ["anno", "data_ultimo", "in_inbox", "n_messaggi", "da", "a", "oggetto", "etichette_attuali",
               "destinazione", "metodo", "dettaglio", "threadId"]
    w2 = wb.create_sheet("Piano")
    w2.append(colonne)
    for c in w2[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for p in piano:
        w2.append([p[c] for c in colonne])
    for i, c in enumerate(colonne, 1):
        w2.column_dimensions[get_column_letter(i)].width = {"oggetto": 50, "da": 35, "a": 35, "destinazione": 35,
                                                            "dettaglio": 30, "etichette_attuali": 30}.get(c, 12)
    w2.freeze_panes = "A2"
    w2.auto_filter.ref = w2.dimensions
    wb.save(PIANO_XLSX)


# ----------------------------------------------------------------------------- --esegui
def cmd_esegui(solo_anni=None, conferma_automatica=False):
    cfg = carica_config()
    if not os.path.exists(PIANO_CSV):
        sys.exit("Manca piano.csv: lanciare prima --piano e controllare il risultato.")
    with open(PIANO_CSV, encoding="utf-8", newline="") as f:
        piano = [p for p in csv.DictReader(f) if p["destinazione"]]
    if solo_anni:
        piano = [p for p in piano if int(p["anno"]) in solo_anni]
    fatti = set()
    if os.path.exists(LOG):
        with open(LOG, encoding="utf-8") as f:
            fatti = {r.rsplit("FATTO thread ", 1)[1].strip() for r in f if "FATTO thread " in r}
    piano = [p for p in piano if p["threadId"] not in fatti]
    togli_inbox = bool(cfg.get("togli_dalla_posta_in_arrivo", True))
    print(f"Conversazioni da etichettare: {len(piano)} (gia' fatte in precedenza: {len(fatti)}); "
          f"togliere dalla Posta in arrivo: {'SI' if togli_inbox else 'NO'}. Nulla viene cancellato.")
    if not piano:
        return
    if not conferma_automatica and input("Scrivere SI per modificare la casella Gmail: ").strip().upper() != "SI":
        print("Annullato: nessuna modifica.")
        return

    svc = servizio()
    etichette = aggiorna_stato_da_server(svc)["etichette"]
    for dest in sorted({p["destinazione"] for p in piano}):
        # Gmail vuole anche le etichette madri: A/B richiede A
        parti = dest.split("/")
        for k in range(1, len(parti) + 1):
            nome = "/".join(parti[:k])
            if nome not in etichette:
                r = con_ritentativi(lambda: svc.users().labels().create(userId="me", body={
                    "name": nome, "labelListVisibility": "labelShow", "messageListVisibility": "show"}).execute())
                etichette[nome] = r["id"]
                log(f"Creata etichetta {nome}")
    gruppi = defaultdict(list)
    for p in piano:
        gruppi[p["destinazione"]].append(p)
    for dest, elenco in gruppi.items():
        corpo = {"addLabelIds": [etichette[dest]], "removeLabelIds": ["INBOX"] if togli_inbox else []}
        for i in range(0, len(elenco), 200):
            blocco = elenco[i:i + 200]
            ids = [m for p in blocco for m in p["messageIds"].split("|")]
            for j in range(0, len(ids), 1000):       # batchModify: max 1000 id per chiamata
                con_ritentativi(lambda: svc.users().messages().batchModify(
                    userId="me", body=dict(corpo, ids=ids[j:j + 1000])).execute())
            with open(LOG, "a", encoding="utf-8") as f:
                for p in blocco:
                    f.write(f"FATTO thread {p['threadId']}\n")
            print(f"\r  {dest:<45} {min(i + 200, len(elenco))}/{len(elenco)}", end="")
            time.sleep(0.3)
        print()
        log(f"{dest}: {len(elenco)} conversazioni etichettate")
    log("Esecuzione completata")


def leggi_anni(testo):
    if not testo:
        return None
    if "-" in testo:
        x, y = testo.split("-")
        return set(range(int(x), int(y) + 1))
    return {int(v) for v in testo.split(",")}


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Riordino di un archivio Gmail con etichette (vedi LEGGIMI.txt)")
    ap.add_argument("--login", action="store_true", help="autorizza l'accesso a Gmail")
    ap.add_argument("--estrai", action="store_true", help="scarica i metadati dei messaggi (sola lettura)")
    ap.add_argument("--query", type=str, help="con --estrai: filtro Gmail (default: query_estrazione del config)")
    ap.add_argument("--piano", action="store_true", help="prova a secco: calcola il piano, nessuna modifica")
    ap.add_argument("--esegui", action="store_true", help="applica il piano alla casella (chiede conferma)")
    ap.add_argument("--anni", type=str, help="con --esegui: solo questi anni, es. 2013-2022 oppure 2023,2024")
    ap.add_argument("--si", action="store_true", help="con --esegui: non chiedere conferma")
    a = ap.parse_args()
    if a.login:
        cmd_login()
    if a.estrai:
        cmd_estrai(a.query or carica_config().get("query_estrazione", "-in:draft"))
    if a.piano:
        cmd_piano()
    if a.esegui:
        cmd_esegui(leggi_anni(a.anni), a.si)
    if not any([a.login, a.estrai, a.piano, a.esegui]):
        ap.print_help()
